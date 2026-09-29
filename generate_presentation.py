import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # completely blank

    # Color Palette (Clean, Professional Developer Style)
    COLOR_BG = RGBColor(0xFA, 0xFA, 0xF9)        # Very light warm gray
    COLOR_CARD = RGBColor(0xFF, 0xFF, 0xFF)      # Pure white
    COLOR_BORDER = RGBColor(0xE2, 0xE8, 0xF0)    # Slate light border
    COLOR_TEXT_MAIN = RGBColor(0x0F, 0x17, 0x2A) # Deep slate almost black
    COLOR_TEXT_SUB = RGBColor(0x33, 0x41, 0x55)  # Slate dark gray
    COLOR_TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)# Muted gray
    COLOR_PRIMARY = RGBColor(0xC2, 0x41, 0x0C)   # Rust orange / Amber dark (Accent)
    COLOR_NAVY = RGBColor(0x1E, 0x29, 0x3B)      # Slate Navy
    COLOR_BLUE = RGBColor(0x02, 0x84, 0xC7)      # Sky blue deep
    COLOR_ACCENT_BG = RGBColor(0xFF, 0xED, 0xD5) # Warm tint

    FONT_HEAD = 'Meiryo'
    FONT_BODY = 'Meiryo'

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background() # no line
        return bg

    def add_header(slide, number_str, category_str, title_str, message_str):
        # Header text frame
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.3))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"{number_str} ｜ {category_str}"
        p0.font.name = FONT_HEAD
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY
        p0.space_after = Pt(2)

        p1 = tf.add_paragraph()
        p1.text = title_str
        p1.font.name = FONT_HEAD
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_NAVY
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = message_str
        p2.font.name = FONT_BODY
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_BLUE

    def add_image_slot(slide, left, top, width, height, title_text, desc_text, image_path=None):
        if image_path and os.path.exists(image_path):
            try:
                pic = slide.shapes.add_picture(image_path, left, top, width, height)
                return pic
            except Exception as e:
                pass

        # Card background for image
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
        box.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
        box.line.width = Pt(1.5)

        # Placeholder text frame
        tf = box.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = f"{title_text}"
        p0.font.name = FONT_HEAD
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_NAVY
        p0.alignment = PP_ALIGN.CENTER
        p0.space_after = Pt(6)

        p1 = tf.add_paragraph()
        p1.text = f"{desc_text}"
        p1.font.name = FONT_BODY
        p1.font.size = Pt(11)
        p1.font.color.rgb = COLOR_TEXT_MUTED
        p1.alignment = PP_ALIGN.CENTER
        return box

    # ═══════════════════════════════════════════════════════════
    # SLIDE 0: 表紙
    # ═══════════════════════════════════════════════════════════
    s0 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s0)

    # Accent top bar
    top_bar = s0.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_PRIMARY
    top_bar.line.fill.background()

    # Title box
    tbox = s0.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.5), Inches(3.5))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "作品プレゼンテーション（5分）"
    p0.font.name = FONT_HEAD
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    p0.space_after = Pt(8)

    p1 = tf.add_paragraph()
    p1.text = "ダッシューティングダック"
    p1.font.name = FONT_HEAD
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY
    p1.space_after = Pt(4)

    p2 = tf.add_paragraph()
    p2.text = "DASH SHOOTING DUCK"
    p2.font.name = 'Arial'
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_after = Pt(14)

    p3 = tf.add_paragraph()
    p3.text = "足音や銃声で敵と駆け引きする、3Dシューティングアクション"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(17)
    p3.font.color.rgb = COLOR_TEXT_SUB
    p3.space_after = Pt(24)

    # Tags
    tags = ["個人制作", "C++ / DirectX 12", "自作エディタ・ツール開発", "敵AI知覚連携", "シューティング手触り設計"]
    tag_x = Inches(1.2)
    for t in tags:
        pill = s0.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tag_x, Inches(4.8), Inches(1.8), Inches(0.45))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        pill.line.color.rgb = COLOR_BORDER
        pill.line.width = Pt(1)
        ptf = pill.text_frame
        ptf.word_wrap = False
        pp = ptf.paragraphs[0]
        pp.text = t
        pp.font.name = FONT_BODY
        pp.font.size = Pt(9.5)
        pp.font.bold = True
        pp.font.color.rgb = COLOR_NAVY
        pp.alignment = PP_ALIGN.CENTER
        tag_x += Inches(1.9)

    # Author
    abox = s0.shapes.add_textbox(Inches(7.5), Inches(5.8), Inches(4.5), Inches(1.0))
    atf = abox.text_frame
    ap0 = atf.paragraphs[0]
    ap0.text = "発表者：コウダ アユ"
    ap0.font.name = FONT_HEAD
    ap0.font.size = Pt(15)
    ap0.font.bold = True
    ap0.font.color.rgb = COLOR_NAVY
    ap0.alignment = PP_ALIGN.RIGHT

    ap1 = atf.add_paragraph()
    ap1.text = "制作期間：約6ヶ月（エンジン＋ゲーム＋ツール）"
    ap1.font.name = FONT_BODY
    ap1.font.size = Pt(12)
    ap1.font.color.rgb = COLOR_TEXT_MUTED
    ap1.alignment = PP_ALIGN.RIGHT

    # ═══════════════════════════════════════════════════════════
    # SLIDE 1: ゲーム概要
    # ═══════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)
    add_header(s1, "01", "ゲーム概要", "ゲーム概要：廃工場を制圧し、物資を持ち帰って生還せよ", "「撃ちまくる爽快感」だけでなく、「逃げるか・挑むか」の判断を作るゲームデザイン")

    # Overview 4 blocks
    specs = [
        ("ジャンル", "3Dシューティングアクション"),
        ("クリア目標", "3箇所の拠点破壊 ➔ 脱出"),
        ("制限時間", "5分（超過でMIA失敗）"),
        ("開発スタイル", "個人制作（プログラム全般）")
    ]
    spec_x = Inches(0.8)
    for label, val in specs:
        box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, spec_x, Inches(1.8), Inches(2.75), Inches(0.7))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = COLOR_BORDER
        box.line.width = Pt(1)
        btf = box.text_frame
        btf.margin_top = Inches(0.08)
        p0 = btf.paragraphs[0]
        p0.text = label
        p0.font.size = Pt(9.5)
        p0.font.color.rgb = COLOR_TEXT_MUTED
        p1 = btf.add_paragraph()
        p1.text = val
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_NAVY
        spec_x += Inches(2.95)

    # Left content box
    cbox = s1.shapes.add_textbox(Inches(0.8), Inches(2.7), Inches(6.0), Inches(4.2))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    p0 = ctf.paragraphs[0]
    p0.text = "● 走る足音や銃声が命取りになるステルス要素"
    p0.font.size = Pt(13.5)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_NAVY
    p0.space_after = Pt(3)

    p1 = ctf.add_paragraph()
    p1.text = "敵は視界だけでなく「音」でもプレイヤーを探知します。不用意に走ったり発砲すると周囲の敵を引き寄せるため、隠れてやり過ごすか、先手を打って仕掛けるかの判断が生まれます。"
    p1.font.size = Pt(11.5)
    p1.font.color.rgb = COLOR_TEXT_SUB
    p1.space_after = Pt(14)

    p2 = ctf.add_paragraph()
    p2.text = "● 脱出直前に訪れる「増援クライマックス」"
    p2.font.size = Pt(13.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_NAVY
    p2.space_after = Pt(3)

    p3 = ctf.add_paragraph()
    p3.text = "3つの拠点を破壊すると脱出ゲートが開きますが、同時にサイレン警報が鳴り敵の増援が急襲。「最後の最後が一番危ない」という起承転結をコードで演出しました。"
    p3.font.size = Pt(11.5)
    p3.font.color.rgb = COLOR_TEXT_SUB
    p3.space_after = Pt(14)

    p4 = ctf.add_paragraph()
    p4.text = "【目指したゲーム性】コミカルなアヒルの見た目とは裏腹に、リソース（弾薬・回復）と位置取りをシビアに管理する本格的なシューター体験を目指しました。"
    p4.font.size = Pt(11)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_PRIMARY

    # Right Image
    add_image_slot(s1, Inches(7.1), Inches(2.7), Inches(5.4), Inches(4.2), "全体ゲームプレイ画面", "トップダウン視点・廃工場マップ・敵との位置関係・アヒル主人公")

    # ═══════════════════════════════════════════════════════════
    # SLIDE 2: 敵AIと戦闘手触り
    # ═══════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "02", "技術のこだわり ①", "敵AIの知覚連携と、撃ち合いの手触り", "ただ突っ込んでくるだけの敵をなくし、音と仲間連携でプレイヤーを追い詰める")

    cbox = s2.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(6.0), Inches(5.0))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    p0 = ctf.paragraphs[0]
    p0.text = "● 視覚だけでなく「音」でプレイヤーを探す敵AI"
    p0.font.size = Pt(13.5)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_NAVY
    p0.space_after = Pt(3)

    p1 = ctf.add_paragraph()
    p1.text = "銃声やダッシュした足音を聞きつけて、その発生場所に警戒しながら集まります。視界（65度）に入っても、物陰に隠れれば見失うため、ステルスでやり過ごす緊張感を作りました。"
    p1.font.size = Pt(11.5)
    p1.font.color.rgb = COLOR_TEXT_SUB
    p1.space_after = Pt(14)

    p2 = ctf.add_paragraph()
    p2.text = "● 仲間を呼んで包囲してくる連携ロジック"
    p2.font.size = Pt(13.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_NAVY
    p2.space_after = Pt(3)

    p3 = ctf.add_paragraph()
    p3.text = "プレイヤーを見つけた敵は周囲の仲間に警戒を伝達し、一斉に集結して挟み撃ちにしてきます。単体で突撃してくるだけの単調さをなくし、常に退路を考えさせます。"
    p3.font.size = Pt(11.5)
    p3.font.color.rgb = COLOR_TEXT_SUB
    p3.space_after = Pt(14)

    p4 = ctf.add_paragraph()
    p4.text = "● 立ち止まって狙うか、逃げながら撃つかのメリハリ"
    p4.font.size = Pt(13.5)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_NAVY
    p4.space_after = Pt(3)

    p5 = ctf.add_paragraph()
    p5.text = "移動しながら撃つと弾が大きくブレる「移動ペナルティ」を実装。足を止めて正確に当てるメリットを強調しました。緊急回避（ローリング）は無敵があるもののスタミナを消費し、連発できない設計にしています。"
    p5.font.size = Pt(11.5)
    p5.font.color.rgb = COLOR_TEXT_SUB
    p5.space_after = Pt(14)

    p6 = ctf.add_paragraph()
    p6.text = "【実装の工夫】敵の数値をただ強くするのではなく、「音を立てずに進む」「見つかったら囲まれる前に引く」というプレイヤーの駆け引きをコードで作りました。"
    p6.font.size = Pt(11)
    p6.font.bold = True
    p6.font.color.rgb = COLOR_PRIMARY

    # Right Image
    img_s2 = os.path.join(os.path.dirname(__file__), "images", "behavior_tree_concept.png")
    add_image_slot(s2, Inches(7.1), Inches(1.9), Inches(5.4), Inches(5.0), "敵AI ビヘイビアツリー意思決定設計図", "視覚・聴覚で索敵し、追跡や攻撃を選択するロジック", image_path=img_s2)

    # ═══════════════════════════════════════════════════════════
    # SLIDE 3: UI演出とゲームループ
    # ═══════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "03", "技術のこだわり ②", "状況が直感的に伝わる情報提示UI", "文字を読まなくても、視覚エフェクトで「危険度」と「被弾方向」がわかる工夫")

    cbox = s3.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(6.0), Inches(5.0))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    p0 = ctf.paragraphs[0]
    p0.text = "● どこまで音が届いたかがわかる「音紋リング」"
    p0.font.size = Pt(13.5)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_NAVY
    p0.space_after = Pt(3)

    p1 = ctf.add_paragraph()
    p1.text = "走ったり発砲した瞬間、「敵にどこまで音が届いているか」を地面に広がる円形エフェクトで可視化。「今のはマズい、走って逃げよう」とプレイヤー自身が直感的に気づけるようにしました。"
    p1.font.size = Pt(11.5)
    p1.font.color.rgb = COLOR_TEXT_SUB
    p1.space_after = Pt(14)

    p2 = ctf.add_paragraph()
    p2.text = "● 死角から撃たれても即座にわかる「360度被弾アロー」"
    p2.font.size = Pt(13.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_NAVY
    p2.space_after = Pt(3)

    p3 = ctf.add_paragraph()
    p3.text = "トップダウン視点でもどこから撃たれたか迷わないよう、プレイヤー中心の円環矢印で弾の飛来方向を明示。体力が減ると画面縁が暗赤色に拍動し、ピンチを体感させます。"
    p3.font.size = Pt(11.5)
    p3.font.color.rgb = COLOR_TEXT_SUB
    p3.space_after = Pt(14)

    p4 = ctf.add_paragraph()
    p4.text = "● 物資探索（死体漁り）の無防備タイマーとリザルト精算"
    p4.font.size = Pt(13.5)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_NAVY
    p4.space_after = Pt(3)

    p5 = ctf.add_paragraph()
    p5.text = "敵の死体を漁るには長押しが必要（0.8秒）。途中で動くとキャンセルされるため、周囲を警戒しながら欲張る駆け引きを表現。クリア時には換金総額や命中率を集計します。"
    p5.font.size = Pt(11.5)
    p5.font.color.rgb = COLOR_TEXT_SUB
    p5.space_after = Pt(14)

    p6 = ctf.add_paragraph()
    p6.text = "【設計の工夫】UI描画クラスがゲームロジックを直接触らないよう、読み取り専用のコンテキスト構造体を受け取る疎結合設計にし、演出の追加や変更を容易にしました。"
    p6.font.size = Pt(11)
    p6.font.bold = True
    p6.font.color.rgb = COLOR_PRIMARY

    # Right Image
    img_s3 = os.path.join(os.path.dirname(__file__), "images", "class_design.png")
    add_image_slot(s3, Inches(7.1), Inches(1.9), Inches(5.4), Inches(5.0), "クラス設計と疎結合アーキテクチャ", "責任分離を両立したノードとコンテキスト設計", image_path=img_s3)

    # ═══════════════════════════════════════════════════════════
    # SLIDE 4: 自作精密衝突判定基盤
    # ═══════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "04", "技術のこだわり ③", "高速戦闘を支える、自作の精密衝突判定基盤", "「弾がすり抜ける」「斜め壁に引っかかる」などの手触りストレスを数学的に徹底排除")

    cbox = s4.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(6.0), Inches(5.0))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    p0 = ctf.paragraphs[0]
    p0.text = "● 高速弾のすり抜けを完全阻止する「Swept-Capsule CCD」"
    p0.font.size = Pt(13.5)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_NAVY
    p0.space_after = Pt(3)

    p1 = ctf.add_paragraph()
    p1.text = "前後フレームの弾丸移動を連続した線分とし、障害物との交差パラメータ t（0.0〜1.0）を解く連続衝突判定（CCD）を自作。どんなに高速な弾丸でも薄い壁をすり抜けない強固な着弾保証を実現しました。"
    p1.font.size = Pt(11.5)
    p1.font.color.rgb = COLOR_TEXT_SUB
    p1.space_after = Pt(14)

    p2 = ctf.add_paragraph()
    p2.text = "● 自由な角度の壁に密着できる「Slab法 任意回転OBB」"
    p2.font.size = Pt(13.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_NAVY
    p2.space_after = Pt(3)

    p3 = ctf.add_paragraph()
    p3.text = "AABB（軸平行）では不可能な斜めの遮蔽物・障害物に対し、3x3ローカル回転基底へ射影するSlab法を採用。傾いたオブジェクトに対しても隙間や引っかかりのない精密な当たり判定を構築しました。"
    p3.font.size = Pt(11.5)
    p3.font.color.rgb = COLOR_TEXT_SUB
    p3.space_after = Pt(14)

    p4 = ctf.add_paragraph()
    p4.text = "● オブジェクト増加でも破綻しない「空間分割ハッシュ O(1)」"
    p4.font.size = Pt(13.5)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_NAVY
    p4.space_after = Pt(3)

    p5 = ctf.add_paragraph()
    p5.text = "総当たり判定 O(N²) を廃止し、2048エントリの空間ハッシュグリッド（セル幅10m）を実装。弾丸近傍のオブジェクトのみを定数時間 O(1) で高速探索し、60fps固定を堅持します。"
    p5.font.size = Pt(11.5)
    p5.font.color.rgb = COLOR_TEXT_SUB
    p5.space_after = Pt(14)

    p6 = ctf.add_paragraph()
    p6.text = "【エンジンの真価】市販エンジンに頼らず、交差数式や空間ハッシュを自らC++でコーディングすることで、ゲームに求められる「当たり判定の心地よさ」と「60fps固定の軽さ」を両立させました。"
    p6.font.size = Pt(11)
    p6.font.bold = True
    p6.font.color.rgb = COLOR_PRIMARY

    # Right Image
    img_s4 = os.path.join(os.path.dirname(__file__), "images", "slide4_collision_fixed.png")
    add_image_slot(s4, Inches(7.1), Inches(1.9), Inches(5.4), Inches(5.0), "自作精密衝突判定システム（Slab法 OBB ＆ Swept CCD）", "空間ハッシュ O(1) 探索・トンネリング（すり抜け）ゼロ保証", image_path=img_s4)

    # ═══════════════════════════════════════════════════════════
    # SLIDE 5: 相談事項
    # ═══════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "05", "社員相談", "現場のプログラマ社員の方に相談したいこと", "作品のこれからの拡張と、就職活動での技術アピールについてアドバイスをいただきたいです")

    # Left Box
    b1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8))
    b1.fill.solid()
    b1.fill.fore_color.rgb = COLOR_CARD
    b1.line.color.rgb = COLOR_BORDER
    b1.line.width = Pt(1.5)

    tf1 = b1.text_frame
    tf1.word_wrap = True
    p0 = tf1.paragraphs[0]
    p0.text = "01 ｜ 作品についての相談"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    p0.space_after = Pt(8)

    p1 = tf1.add_paragraph()
    p1.text = "敵AIやゲーム規模が大きくなったときの\n現場のコード設計について"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY
    p1.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "本作ではFSMに加え、ImGuiでノードを繋ぐ「自作ビヘイビアツリーエディタ」を構築しました。\n\n今後、敵の種類（近接兵・スナイパー等）や複雑な包囲行動を追加していくにあたり、\n\n商業開発の現場では、「AIの肥大化を防ぎ、プランナーとも協調しながら破綻せずに拡張していくために、どのようなBehavior Treeやタスクシステムの設計・運用をしているか」を伺いたいです。"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_TEXT_SUB

    # Right Box
    b2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.8))
    b2.fill.solid()
    b2.fill.fore_color.rgb = COLOR_CARD
    b2.line.color.rgb = COLOR_BORDER
    b2.line.width = Pt(1.5)

    tf2 = b2.text_frame
    tf2.word_wrap = True
    p0 = tf2.paragraphs[0]
    p0.text = "02 ｜ 就職活動全般の相談"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    p0.space_after = Pt(8)

    p1 = tf2.add_paragraph()
    p1.text = "自作ツール・エンジン経験と\nゲームプレイ実装のアピールバランス"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY
    p1.space_after = Pt(14)

    p2 = tf2.add_paragraph()
    p2.text = "自作DirectX 12エンジンや開発用エディタを作ることで制作環境を整えつつ、本作では「遊んで気持ちいいゲームプレイ」を最優先に作り込みました。\n\n商用エンジン（Unreal/Unity）での開発が主流の今、新卒採用の現場において、\n\n「ツールやエンジンを自作できる技術力」と「ゲームの手触りや面白さを作る力」をどのような切り口や比率で伝えると、ゲームプログラマとして最も評価されるかを教えていただきたいです。"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_TEXT_SUB

    # Save presentation
    import shutil
    base_dir = os.path.dirname(os.path.abspath(__file__))
    out_ascii = os.path.join(base_dir, "DashShootingDuck_Presentation.pptx")
    out_jp = os.path.join(base_dir, "ダッシューティングダック_作品プレゼン資料.pptx")
    prs.save(out_ascii)
    try:
        shutil.copyfile(out_ascii, out_jp)
    except Exception:
        pass
    print(f"Presentation saved successfully to: {out_ascii}")

if __name__ == "__main__":
    create_deck()
