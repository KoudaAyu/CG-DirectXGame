import os
import shutil
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

def main():
    base_dir = r"c:\Users\3329a\OneDrive\デスクトップ\Engine"
    preview_dir = os.path.join(base_dir, "slides_preview")
    output_crop_dir = os.path.join(base_dir, "DashShootingDuck_Slides")
    downloads_dir = r"C:\Users\3329a\Downloads"
    downloads_slides_dir = os.path.join(downloads_dir, "DashShootingDuck_Slides")

    os.makedirs(output_crop_dir, exist_ok=True)
    os.makedirs(downloads_slides_dir, exist_ok=True)

    # 1. Bounding box for exact 16:9 presentation frame: (290, 99, 2089, 1111)
    # To be extremely clean and avoid border antialiasing, box is (290, 99, 2089, 1111)
    crop_box = (290, 99, 2089, 1111)
    cropped_images = []

    print("--- 1. Cropping pure 16:9 slides ---")
    for i in range(1, 7):
        src_path = os.path.join(preview_dir, f"slide_{i}.png")
        if not os.path.exists(src_path):
            raise FileNotFoundError(f"Missing preview slide: {src_path}")

        im = Image.open(src_path).convert("RGB")
        cropped = im.crop(crop_box)
        
        dst_path = os.path.join(output_crop_dir, f"slide_{i}.png")
        cropped.save(dst_path, "PNG", quality=95)
        cropped_images.append(dst_path)
        print(f"Saved: {dst_path} (Size: {cropped.size})")

        # Copy to downloads
        dst_dl = os.path.join(downloads_slides_dir, f"slide_{i}.png")
        shutil.copy2(dst_path, dst_dl)

    # 2. Build 16:9 PowerPoint Presentation
    print("\n--- 2. Building 16:9 PowerPoint Presentation ---")
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    for i, img_path in enumerate(cropped_images, 1):
        slide = prs.slides.add_slide(blank_layout)
        # Add full bleed 16:9 image
        slide.shapes.add_picture(
            img_path,
            Inches(0),
            Inches(0),
            width=prs.slide_width,
            height=prs.slide_height
        )
        print(f"Added Slide {i} to PPTX")

    pptx_path = os.path.join(base_dir, "DashShootingDuck_Presentation.pptx")
    prs.save(pptx_path)
    print(f"\nPPTX saved to: {pptx_path}")

    # Copy PPTX to Downloads
    dl_pptx_path = os.path.join(downloads_dir, "DashShootingDuck_Presentation.pptx")
    shutil.copy2(pptx_path, dl_pptx_path)
    print(f"PPTX copied to Downloads: {dl_pptx_path}")

    # Copy HTML to Downloads
    html_src = os.path.join(base_dir, "college_slides.html")
    html_dl = os.path.join(downloads_dir, "college_slides.html")
    shutil.copy2(html_src, html_dl)
    print(f"HTML copied to Downloads: {html_dl}")

    print("\n[SUCCESS] All files generated and synchronized with Downloads folder!")

if __name__ == "__main__":
    main()
