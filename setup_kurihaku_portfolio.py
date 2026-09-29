import os
import shutil
import pypdf
import openpyxl
from openpyxl.styles import Alignment, Font
import zipfile

downloads = r"C:\Users\3329a\Downloads"
desktop_engine = r"c:\Users\3329a\OneDrive\デスクトップ\Engine"
source_root = r"C:\Users\3329a\Downloads\月次制作\月次制作"
target_root = r"C:\Users\3329a\Downloads\月次制作\クリ博提出用（3作品）"

print("Target folder:", target_root)
os.makedirs(target_root, exist_ok=True)

# 1. 自己PRシート（公田亜優）.xlsx
print("Step 1: Copying and polishing 自己PRシート（公田亜優）.xlsx...")
pr_src = os.path.join(downloads, "公田亜優（日本工学院専門学校_デザインカレッジゲームクリエイター科目四年制）_自己PRシートフォーマット.xlsx")
if not os.path.exists(pr_src):
    for f in os.listdir(downloads):
        if f.endswith(".xlsx") and "自己PRシート" in f:
            pr_src = os.path.join(downloads, f)
            break

wb = openpyxl.load_workbook(pr_src)
ws = wb.active

# 卒業見込みの補完
if not ws['C16'].value:
    ws['C16'] = '2028'
if not ws['F16'].value:
    ws['F16'] = '3'

# 1行整理
ws['I22'] = 'C++(1年/自作エンジン,STL,低レイヤ), DirectX12(1年/パイプライン,CS,最適化), HLSL, Python, Unity, Blender(独自エディタ連携), 2Dツール(アイビス5年/UI制作), Excel, PPT'
ws['I22'].alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
ws['I22'].font = Font(name='游ゴシック', size=9.0)

ws['I23'] = '【趣味】アクション/FPS研究（敵AI挙動・手触り・演出分析）、イラスト制作 【特技】プロファイラによるボトルネック特定、DCC連携ツール開発'
ws['I23'].alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
ws['I23'].font = Font(name='游ゴシック', size=9.0)

pr_dst = os.path.join(target_root, "自己PRシート（公田亜優）.xlsx")
wb.save(pr_dst)
print("Saved:", pr_dst)

# 2. 作品①：ダッシューティングダック
print("\nStep 2: Setting up ダッシューティングダック...")
duck_dir = os.path.join(target_root, "ダッシューティングダック")
os.makedirs(duck_dir, exist_ok=True)

duck_exe_dir = os.path.join(duck_dir, "実行ファイル（ダッシューティングダック）")
duck_src_proj = os.path.join(source_root, "個人制作", "プロジェクト")

def copy_clean_project(src, dst):
    os.makedirs(dst, exist_ok=True)
    for item in os.listdir(src):
        if item in [".vs", ".git"]:
            continue
        s = os.path.join(src, item)
        d = os.path.join(dst, item)
        if os.path.isdir(s):
            if item.lower() in ["intermediate", ".vs"]:
                continue
            copy_clean_project(s, d)
        else:
            ext = os.path.splitext(item)[1].lower()
            if ext in [".obj", ".tlog", ".ipch", ".idb", ".ilk", ".iobj"]:
                continue
            shutil.copy2(s, d)

print("Copying cleaned project for ダッシューティングダック...")
copy_clean_project(duck_src_proj, duck_exe_dir)

# プログラム説明資料（ダッシューティングダック）.pdf
duck_pdf_src = os.path.join(downloads, "LE3B_09_コウダ_アユ_「射撃と遮蔽物の動的物理インタラクション」 の実装.pdf")
if not os.path.exists(duck_pdf_src):
    duck_pdf_src = os.path.join(desktop_engine, "LE3B_09_コウダ_アユ_「射撃と遮蔽物の動的物理インタラクション」 の実装.pdf")
if not os.path.exists(duck_pdf_src):
    duck_pdf_src = os.path.join(source_root, "LE3B_09_コウダ_アユ_プログラミング説明資料.pdf")

duck_pdf_dst = os.path.join(duck_dir, "プログラム説明資料（ダッシューティングダック）.pdf")
shutil.copy2(duck_pdf_src, duck_pdf_dst)
print("Copied duck pdf")

# 作品実演動画（ダッシューティングダック）.mp4
duck_mp4_src = os.path.join(source_root, "個人制作", "作品動画", "個人制作.mp4")
duck_mp4_dst = os.path.join(duck_dir, "作品実演動画（ダッシューティングダック）.mp4")
shutil.copy2(duck_mp4_src, duck_mp4_dst)
print("Copied duck mp4")

# 3. 作品②：雨宿り
print("\nStep 3: Setting up 雨宿り...")
rain_dir = os.path.join(target_root, "雨宿り")
os.makedirs(rain_dir, exist_ok=True)

rain_exe_dir = os.path.join(rain_dir, "実行ファイル（雨宿り）")
os.makedirs(rain_exe_dir, exist_ok=True)

rain_built_src = os.path.join(source_root, "チーム制作", "雨宿り", "ビルド済み作品")
if os.path.exists(rain_built_src):
    for item in os.listdir(rain_built_src):
        s = os.path.join(rain_built_src, item)
        d = os.path.join(rain_exe_dir, item)
        if os.path.isdir(s):
            shutil.copytree(s, d, dirs_exist_ok=True)
        else:
            shutil.copy2(s, d)

rain_code_src = os.path.join(source_root, "チーム制作", "雨宿り", "担当コード")
if os.path.exists(rain_code_src):
    shutil.copytree(rain_code_src, os.path.join(rain_exe_dir, "担当ソースコード"), dirs_exist_ok=True)

rain_pdf_src = os.path.join(downloads, "LE3B_09_コウダ_アユ_雨宿り_プログラム.pdf")
if not os.path.exists(rain_pdf_src):
    rain_pdf_src = os.path.join(source_root, "LE3B_コウダ_アユ_ポートフォリオ.pdf")
rain_pdf_dst = os.path.join(rain_dir, "プログラム説明資料（雨宿り）.pdf")
shutil.copy2(rain_pdf_src, rain_pdf_dst)
print("Copied rain pdf")

rain_mp4_src = os.path.join(source_root, "チーム制作", "雨宿り", "作品動画", "雨宿り.mp4")
rain_mp4_dst = os.path.join(rain_dir, "作品実演動画（雨宿り）.mp4")
shutil.copy2(rain_mp4_src, rain_mp4_dst)
print("Copied rain mp4")

# 4. 作品③：のびへび
print("\nStep 4: Setting up のびへび...")
snake_dir = os.path.join(target_root, "のびへび")
os.makedirs(snake_dir, exist_ok=True)

snake_exe_dir = os.path.join(snake_dir, "実行ファイル（のびへび）")
os.makedirs(snake_exe_dir, exist_ok=True)

snake_built_src = os.path.join(source_root, "チーム制作", "のびへび", "ビルド済み作品")
if os.path.exists(snake_built_src):
    for item in os.listdir(snake_built_src):
        s = os.path.join(snake_built_src, item)
        d = os.path.join(snake_exe_dir, item)
        if os.path.isdir(s):
            shutil.copytree(s, d, dirs_exist_ok=True)
        else:
            shutil.copy2(s, d)

snake_code_src = os.path.join(source_root, "チーム制作", "のびへび", "担当ソースコード")
if os.path.exists(snake_code_src):
    shutil.copytree(snake_code_src, os.path.join(snake_exe_dir, "担当ソースコード"), dirs_exist_ok=True)

portfolio_pdf = os.path.join(downloads, "LE3B_09_コウダ_アユ_ポートフォリオ.pdf")
snake_pdf_dst = os.path.join(snake_dir, "プログラム説明資料（のびへび）.pdf")

if os.path.exists(portfolio_pdf):
    reader = pypdf.PdfReader(portfolio_pdf)
    writer = pypdf.PdfWriter()
    for p_idx in [0, 1, 7, 8]:
        if p_idx < len(reader.pages):
            writer.add_page(reader.pages[p_idx])
    with open(snake_pdf_dst, "wb") as f_out:
        writer.write(f_out)
    print("Created プログラム説明資料（のびへび）.pdf")
else:
    shutil.copy2(os.path.join(source_root, "LE3B_コウダ_アユ_ポートフォリオ.pdf"), snake_pdf_dst)

snake_mp4_src = os.path.join(source_root, "チーム制作", "のびへび", "作品動画", "のびへび.mp4")
snake_mp4_dst = os.path.join(snake_dir, "作品実演動画（のびへび）.mp4")
shutil.copy2(snake_mp4_src, snake_mp4_dst)
print("Copied snake mp4")

print("\nSetup completed successfully!")
